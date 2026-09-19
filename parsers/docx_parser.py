import os
import re
from datetime import datetime
from typing import Dict, Any, List, Optional
from parsers.base import BaseParser

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


ROLE_ALIASES = {
    "user": "user",
    "human": "user",
    "you": "user",
    "me": "user",
    "assistant": "assistant",
    "ai": "assistant",
    "bot": "assistant",
    "deepseek": "assistant",
    "chatgpt": "assistant",
    "claude": "assistant",
    "gemini": "assistant",
    "system": "system",
}

HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
BOLD_RE = re.compile(r"^\*\*(.+?):?\*\*\s*$")
ITALIC_RE = re.compile(r"^\*(.+?):?\*\s*$")


class DocxChatParser(BaseParser):
    """Parser untuk file .docx dengan format chat mentah."""

    def __init__(self, file_path: str):
        super().__init__(file_path)
        self.paragraphs: List[str] = []

    def validate(self) -> bool:
        if not DOCX_AVAILABLE:
            return False
        if not os.path.exists(self.file_path):
            return False
        try:
            doc = Document(self.file_path)
            self.paragraphs = [p.text for p in doc.paragraphs]
            return len(self.paragraphs) > 0
        except Exception:
            return False

    def _match_role(self, candidate: str) -> Optional[str]:
        if not candidate:
            return None
        clean = candidate.strip().strip(":").strip().lower()
        clean = re.sub(r"[^a-z]", "", clean)
        return ROLE_ALIASES.get(clean)

    def _detect_header(self, line: str) -> Optional[str]:
        stripped = line.strip()
        if not stripped:
            return None

        m = HEADING_RE.match(stripped)
        if m:
            role = self._match_role(m.group(2))
            if role:
                return role

        m = BOLD_RE.match(stripped)
        if m:
            role = self._match_role(m.group(1))
            if role:
                return role

        m = ITALIC_RE.match(stripped)
        if m:
            role = self._match_role(m.group(1))
            if role:
                return role

        return None

    def _extract_title(self) -> str:
        for line in self.paragraphs:
            s = line.strip()
            if s.startswith("# ") and not s.startswith("## "):
                return s[2:].strip()
        base = os.path.splitext(os.path.basename(self.file_path))[0]
        return base or "DOCX Chat"

    def list_conversations(self) -> List[Dict[str, Any]]:
        return [{
            "index": 0,
            "title": self._extract_title(),
            "created_at": "",
            "node_count": 0,
        }]

    def parse(self, convo_index: Optional[int] = None) -> Dict[str, Any]:
        if not self.validate():
            raise ValueError(f"Gagal membaca file DOCX: {self.file_path}")

        title = self._extract_title()
        messages: List[Dict[str, str]] = []
        current_role: Optional[str] = None
        current_buffer: List[str] = []

        def flush():
            if current_role is None:
                return
            content = "\n".join(current_buffer).strip()
            if content:
                messages.append({
                    "role": current_role,
                    "content": content,
                    "timestamp": "",
                })

        for line in self.paragraphs:
            role = self._detect_header(line)
            if role:
                flush()
                current_role = role
                current_buffer = []
                continue
            if current_role is not None:
                current_buffer.append(line)

        flush()

        created_at = datetime.fromtimestamp(
            os.path.getmtime(self.file_path)
        ).strftime("%Y-%m-%d %H:%M") if os.path.exists(self.file_path) else ""

        return {
            "title": title,
            "created_at": created_at,
            "messages": messages,
        }